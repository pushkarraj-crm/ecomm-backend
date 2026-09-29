import hashlib
import math
import os
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.utils.text import slugify
from rest_framework.exceptions import ValidationError

from .models import Category, Product, ProductVariant


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 1000
MAX_PRODUCTS = 100
REQUIRED_COLUMNS = {
    'product_key', 'product_name', 'description', 'category_id',
    'size', 'color', 'price', 'stock',
}


def create_bulk_catalog_template():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Catalog'
    sheet.append([
        'product_key', 'product_name', 'description', 'category_id',
        'size', 'color', 'price', 'stock',
    ])
    sheet.append([
        'TSHIRT-001', 'Cotton T-shirt', 'Cotton crew-neck T-shirt',
        1, 'M', 'Blue', 499.00, 25,
    ])
    sheet.append([
        'TSHIRT-001', 'Cotton T-shirt', 'Cotton crew-neck T-shirt',
        1, 'L', 'Blue', 499.00, 18,
    ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:H1'
    instructions = workbook.create_sheet('Instructions')
    instructions.append(['Bulk catalog upload'])
    instructions.append(['Use the Catalog sheet; repeat product_key for each variant.'])
    instructions.append(['Replace category_id 1 with an ID returned by GET /api/products/categories/.'])
    instructions.append(['Products must have matching name, description, and category_id for each product_key.'])
    instructions.append(['Maximum 10 MB, 1000 data rows, and 100 products per workbook.'])
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="seller-catalog-template.xlsx"'
    response['Cache-Control'] = 'private, no-store'
    return response


def _cell_text(value, row, column, required=True):
    if value is None:
        value = ''
    value = str(value).strip()
    if required and not value:
        raise ValidationError({'file': f'Row {row}: {column} is required.'})
    return value


def _read_catalog_rows(upload):
    if upload.size > MAX_FILE_BYTES:
        raise ValidationError({'file': 'Excel file must be 10 MB or smaller.'})
    if os.path.splitext(upload.name or '')[1].lower() != '.xlsx':
        raise ValidationError({'file': 'Upload an .xlsx Excel file.'})
    workbook = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        if sheet.max_row and sheet.max_row > MAX_DATA_ROWS + 1:
            raise ValidationError({'file': f'Excel file cannot exceed {MAX_DATA_ROWS} data rows.'})
        if sheet.max_column and sheet.max_column > 64:
            raise ValidationError({'file': 'Excel file cannot have more than 64 columns.'})
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValidationError({'file': 'The workbook must have a header row.'})
        headers = list(headers)
        while headers and not str(headers[-1] or '').strip():
            headers.pop()
        normalized_headers = [str(value or '').strip().lower() for value in headers]
        populated_headers = [value for value in normalized_headers if value]
        if len(populated_headers) != len(set(populated_headers)):
            raise ValidationError({'file': 'The header row contains duplicate columns.'})
        missing = REQUIRED_COLUMNS - set(normalized_headers)
        if missing:
            raise ValidationError({
                'file': f"Missing columns: {', '.join(sorted(missing))}.",
            })
        indexes = {name: normalized_headers.index(name) for name in REQUIRED_COLUMNS}
        grouped = OrderedDict()
        row_count = 0
        for row_number, raw_row in enumerate(rows, start=2):
            if not raw_row or all(value is None or str(value).strip() == '' for value in raw_row):
                continue
            row_count += 1
            if row_count > MAX_DATA_ROWS:
                raise ValidationError({'file': f'Excel file cannot exceed {MAX_DATA_ROWS} data rows.'})
            cells = {
                column: raw_row[index] if index < len(raw_row) else None
                for column, index in indexes.items()
            }
            key = _cell_text(cells['product_key'], row_number, 'product_key')
            if len(key) > 80:
                raise ValidationError({'file': f'Row {row_number}: product_key must be 80 characters or fewer.'})
            name = _cell_text(cells['product_name'], row_number, 'product_name')
            description = _cell_text(cells['description'], row_number, 'description')
            if len(name) > 255:
                raise ValidationError({'file': f'Row {row_number}: product_name must be 255 characters or fewer.'})
            try:
                category_value = Decimal(str(cells['category_id']))
                price = Decimal(str(cells['price']))
                stock_value = Decimal(str(cells['stock']))
            except (TypeError, ValueError, InvalidOperation):
                raise ValidationError({
                    'file': f'Row {row_number}: category_id, price, or stock is not a valid number.',
                })
            if (
                not category_value.is_finite()
                or category_value != category_value.to_integral_value()
                or not stock_value.is_finite()
                or stock_value != stock_value.to_integral_value()
            ):
                raise ValidationError({
                    'file': f'Row {row_number}: category_id and stock must be whole numbers.',
                })
            category_id = int(category_value)
            stock = int(stock_value)
            if category_id < 1:
                raise ValidationError({'file': f'Row {row_number}: category_id must be positive.'})
            if not price.is_finite() or price <= 0:
                raise ValidationError({'file': f'Row {row_number}: price must be greater than zero.'})
            if stock < 0:
                raise ValidationError({'file': f'Row {row_number}: stock cannot be negative.'})
            if stock > 2147483647:
                raise ValidationError({'file': f'Row {row_number}: stock exceeds the supported maximum.'})
            size = _cell_text(cells['size'], row_number, 'size')
            color = _cell_text(cells['color'], row_number, 'color')
            if len(size) > 50 or len(color) > 50:
                raise ValidationError({'file': f'Row {row_number}: size and color must be 50 characters or fewer.'})
            try:
                price_float = float(price)
            except (OverflowError, ValueError):
                raise ValidationError({'file': f'Row {row_number}: price is too large.'})
            if not math.isfinite(price_float):
                raise ValidationError({'file': f'Row {row_number}: price is too large.'})
            product = grouped.get(key)
            if product is None:
                if len(grouped) >= MAX_PRODUCTS:
                    raise ValidationError({'file': f'Excel file cannot contain more than {MAX_PRODUCTS} products.'})
                product = {
                    'name': name,
                    'description': description,
                    'category_id': category_id,
                    'variants': [],
                    'variant_keys': set(),
                    'first_row': row_number,
                }
                grouped[key] = product
            elif (
                product['name'] != name
                or product['description'] != description
                or product['category_id'] != category_id
            ):
                raise ValidationError({
                    'file': f'Row {row_number}: product details for product_key {key!r} must match its first row.',
                })
            variant_key = (size.casefold(), color.casefold())
            if variant_key in product['variant_keys']:
                raise ValidationError({
                    'file': f'Row {row_number}: duplicate size and color for product_key {key!r}.',
                })
            product['variant_keys'].add(variant_key)
            product['variants'].append({
                'size': size,
                'color': color,
                'price': price_float,
                'stock': stock,
            })
        workbook.close()
    except ValidationError:
        if workbook is not None:
            workbook.close()
        raise
    except Exception as error:
        if workbook is not None:
            workbook.close()
        raise ValidationError({'file': 'Could not read this Excel workbook.'}) from error

    if not grouped:
        raise ValidationError({'file': 'The workbook does not contain any catalog rows.'})
    return grouped


def import_bulk_catalog(seller, upload):
    grouped = _read_catalog_rows(upload)
    return _persist_bulk_catalog(seller, grouped)


@transaction.atomic
def _persist_bulk_catalog(seller, grouped):
    category_ids = {row['category_id'] for row in grouped.values()}
    categories = Category.objects.in_bulk(category_ids)
    unknown_categories = category_ids - set(categories)
    if unknown_categories:
        raise ValidationError({
            'file': f"Unknown category_id values: {', '.join(map(str, sorted(unknown_categories)))}.",
        })

    products = []
    product_keys = list(grouped.items())
    proposed_slugs = []
    for product_key, data in product_keys:
        base = slugify(data['name'])[:155] or 'product'
        key_hash = hashlib.sha256(f'{seller.pk}:{product_key}'.encode('utf-8')).hexdigest()[:12]
        proposed_slugs.append(f'{base}-s{seller.pk}-{key_hash}')
    occupied = set(Product.objects.filter(slug__in=proposed_slugs).values_list('slug', flat=True))

    for (product_key, data), proposed_slug in zip(product_keys, proposed_slugs):
        unique_slug = proposed_slug
        suffix = 1
        while unique_slug in occupied:
            unique_slug = f'{proposed_slug[:175]}-{suffix}'
            suffix += 1
        occupied.add(unique_slug)
        products.append(Product(
            seller=seller,
            category_id=data['category_id'],
            name=data['name'],
            description=data['description'],
            slug=unique_slug,
            status='pending',
        ))

    Product.objects.bulk_create(products)
    variants = []
    for product, (_, data) in zip(products, product_keys):
        variants.extend(
            ProductVariant(product=product, **variant_data)
            for variant_data in data['variants']
        )
    ProductVariant.objects.bulk_create(variants, batch_size=500)
    return products, len(variants)
